from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse_lazy
from django.contrib import messages
from django.shortcuts import redirect, render
from django.views.generic import CreateView, DeleteView, ListView, UpdateView
from django.views import View

from .forms import (
    CategoriaForm,
    ProductoForm,
    ProveedorForm,
    StockMoverProductoBodegaForm,
    RestarProductoForm,
)
from .models import (
    Categoria,
    Producto,
    Proveedor,
    Stock,
    StockActual,
    StockMovimiento,
    StockMotivo,
)
from .filters import StockFilter, StockActualAdminFilter
from django_filters.views import FilterView

from apps.notificaciones.models import Notificaciones, EstadoNotificacion

import openpyxl
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.utils.decorators import method_decorator
from django.views import View

from .models import Producto


class StaffSuperuserRequiredMixin(UserPassesTestMixin):
    def test_func(self):
        return self.request.user.is_staff or self.request.user.is_superuser


# ==================================================================================
# LISTAR PROVEEDORES
# ==================================================================================
class ProveedorListView(LoginRequiredMixin, StaffSuperuserRequiredMixin, ListView):
    model = Proveedor
    paginate_by = 10
    template_name = "inventario/proveedor/proveedor_list.html"

    def get_queryset(self):
        # Opcional: mostrar todos los proveedores o filtrar
        return Proveedor.objects.all()


# ==================================================================================
# CREAR PROVEEDORES
# ==================================================================================
class ProveedorCreateView(LoginRequiredMixin, StaffSuperuserRequiredMixin, CreateView):
    model = Proveedor
    form_class = ProveedorForm
    template_name = "inventario/proveedor/proveedor_form.html"
    success_url = reverse_lazy("proveedor_list")

    def form_valid(self, form):
        form.instance.usuario = self.request.user
        return super().form_valid(form)


# ==================================================================================
# ACTUALIZAR PROVEEDORES
# ==================================================================================
class ProveedorUpdateView(LoginRequiredMixin, StaffSuperuserRequiredMixin, UpdateView):
    model = Proveedor
    form_class = ProveedorForm
    template_name = "inventario/proveedor/proveedor_form.html"
    success_url = reverse_lazy("proveedor_list")


# ==================================================================================
# ELIMINAR PROVEEDORES
# ==================================================================================
class ProveedorDeleteView(LoginRequiredMixin, StaffSuperuserRequiredMixin, DeleteView):
    model = Proveedor
    template_name = "inventario/proveedor/proveedor_confirm_delete.html"
    success_url = reverse_lazy("proveedor_list")


# ===========================================================================================
# LISTAR CATEGORÍAS
# ===========================================================================================
class CategoriaListView(LoginRequiredMixin, StaffSuperuserRequiredMixin, ListView):
    model = Categoria
    paginate_by = 10
    template_name = "inventario/categoria/categoria_list.html"
    context_object_name = "categorias"

    def get_queryset(self):
        return Categoria.objects.all().order_by("-created")


# ===========================================================================================
# CREAR CATEGORÍA
# ===========================================================================================
class CategoriaCreateView(LoginRequiredMixin, StaffSuperuserRequiredMixin, CreateView):
    model = Categoria
    form_class = CategoriaForm
    template_name = "inventario/categoria/categoria_form.html"
    success_url = reverse_lazy("categoria_list")

    def form_valid(self, form):
        form.instance.usuario = self.request.user
        return super().form_valid(form)


# ===========================================================================================
# EDITAR CATEGORÍA
# ===========================================================================================
class CategoriaUpdateView(LoginRequiredMixin, StaffSuperuserRequiredMixin, UpdateView):
    model = Categoria
    form_class = CategoriaForm
    template_name = "inventario/categoria/categoria_form.html"
    success_url = reverse_lazy("categoria_list")


# ===========================================================================================
# ELIMINAR CATEGORÍA
# ===========================================================================================
class CategoriaDeleteView(LoginRequiredMixin, StaffSuperuserRequiredMixin, DeleteView):
    model = Categoria
    template_name = "inventario/categoria/categoria_confirm_delete.html"
    success_url = reverse_lazy("categoria_list")


# ===========================================================================================
# LISTAR PRODUCTOS
# ===========================================================================================
class ProductoListView(LoginRequiredMixin, StaffSuperuserRequiredMixin, ListView):
    model = Producto
    paginate_by = 10
    template_name = "inventario/producto/producto_list.html"
    context_object_name = "productos"

    def get_queryset(self):
        return Producto.objects.select_related("proveedor", "categoria").order_by(
            "-created"
        )


# ===========================================================================================
# CREAR PRODUCTO
# ===========================================================================================
class ProductoCreateView(LoginRequiredMixin, StaffSuperuserRequiredMixin, CreateView):
    model = Producto
    form_class = ProductoForm
    template_name = "inventario/producto/producto_form.html"
    success_url = reverse_lazy("producto_list")

    def form_valid(self, form):
        form.instance.usuario = self.request.user

        response = super().form_valid(form)

        Notificaciones.objects.create(
            usuario=self.request.user,
            estado=EstadoNotificacion.EXITO,
            titulo="Producto creado",
            mensaje="Producto creado exitosamente por un administrador.",
        )

        return response


# ===========================================================================================
# EDITAR PRODUCTO
# ===========================================================================================
class ProductoUpdateView(LoginRequiredMixin, StaffSuperuserRequiredMixin, UpdateView):
    model = Producto
    form_class = ProductoForm
    template_name = "inventario/producto/producto_form.html"
    success_url = reverse_lazy("producto_list")

    def form_valid(self, form):
        form.instance.usuario = self.request.user

        response = super().form_valid(form)

        Notificaciones.objects.create(
            usuario=self.request.user,
            estado=EstadoNotificacion.ADVERTENCIA,
            titulo="Producto modificado",
            mensaje="Producto modificado por un administrador.",
        )

        return response


# ===========================================================================================
# ELIMINAR PRODUCTO
# ===========================================================================================
class ProductoDeleteView(LoginRequiredMixin, StaffSuperuserRequiredMixin, DeleteView):
    model = Producto
    template_name = "inventario/producto/producto_confirm_delete.html"
    success_url = reverse_lazy("producto_list")

    def form_valid(self, form):

        response = super().form_valid(form)

        Notificaciones.objects.create(
            usuario=self.request.user,
            estado=EstadoNotificacion.ADVERTENCIA,
            titulo="Producto eliminado",
            mensaje="Producto eliminado por un administrador.",
        )
        return response


# ===========================================================================================
# LISTAR STOCK
# ===========================================================================================
class StockListView(LoginRequiredMixin, StaffSuperuserRequiredMixin, ListView):
    model = Stock
    paginate_by = 10
    template_name = "inventario/stock/stock_list.html"
    context_object_name = "stocks"

    def get_queryset(self):
        qs = Stock.objects.select_related(
            "usuario", "producto", "usuario_origen", "usuario_destino"
        ).order_by("-created")
        self.filterset = StockFilter(self.request.GET, queryset=qs)
        return self.filterset.qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["filter"] = self.filterset
        return context


# ===========================================================================================
# MOVER PRODUCTO STOCK
# ===========================================================================================
class MoverProductoBodegaView(LoginRequiredMixin, CreateView):
    model = Stock
    form_class = StockMoverProductoBodegaForm
    template_name = "inventario/stock/stock_form.html"
    success_url = reverse_lazy("stock_list")

    def form_valid(self, form):
        form.instance.usuario = self.request.user
        return super().form_valid(form)


# ===========================================================================================
# MI STOCK
# ===========================================================================================
class MiStockListView(LoginRequiredMixin, ListView):
    model = StockActual
    paginate_by = 10
    template_name = "inventario/stock/mi_stock.html"
    context_object_name = "stocks"

    def get_queryset(self):
        return (
            StockActual.objects.filter(usuario=self.request.user)
            .select_related("usuario", "producto")
            .order_by("producto__nombre")
        )

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # 🔥 Agregar historial de restas (movimientos tipo SALIDA)
        context["historial_restas"] = (
            Stock.objects.filter(usuario=self.request.user, movimiento="SALIDA")
            .select_related("producto")
            .order_by("-created")  # created viene del TimeStampedModel
        )

        return context


# ===============================================================================================
# RESTAR PRODUCTO VIEW
# ===============================================================================================


class RestarProductoView(View):
    template_name = "inventario/stock/restar_producto.html"

    def get(self, request):
        form = RestarProductoForm(usuario=request.user)
        return render(request, self.template_name, {"form": form})

    def post(self, request):
        form = RestarProductoForm(request.POST, usuario=request.user)

        if form.is_valid():
            producto_stock = form.cleaned_data["producto"]
            cantidad = form.cleaned_data["cantidad"]
            motivo = form.cleaned_data["motivo"]  # nuevo

            # Validar stock disponible
            if producto_stock.cantidad < cantidad:
                form.add_error("cantidad", "No tienes suficiente stock disponible.")
                return render(request, self.template_name, {"form": form})

            # Crear movimiento de SALIDA
            Stock.objects.create(
                usuario=request.user,
                producto=producto_stock.producto,
                usuario_origen=request.user,
                cantidad_movida=cantidad,
                movimiento=StockMovimiento.Salida,
                motivo=StockMotivo.Ticket,  # motivo general
                motivo_personalizado=motivo or "",  # motivo escrito por el usuario
            )

            return redirect("mi_stock")

        return render(request, self.template_name, {"form": form})


# ===============================================================================================
# HISTORIAL DE RESTAS (ADMIN)
# ===============================================================================================


class HistorialRestasAdministradorView(LoginRequiredMixin, FilterView):
    model = Stock
    template_name = "inventario/admin/historial_restas.html"
    context_object_name = "restas"
    filterset_class = StockFilter
    paginate_by = 20

    def get_queryset(self):
        return (
            Stock.objects.filter(movimiento="SALIDA")  # tus restas
            .select_related("producto", "usuario")
            .order_by("-created")
        )


# ===============================================================================================
# ELIMINAR UNA RESTA - ADMIN
# ===============================================================================================


class EliminarRestaView(LoginRequiredMixin, View):
    def get(self, request, pk):

        # Buscar movimiento SALIDA
        mov = Stock.objects.filter(id=pk, movimiento="SALIDA").first()
        if not mov:
            return redirect("historial_restas_admin")

        # Obtener o crear StockActual del usuario que tenía el producto
        stock_actual, creado = StockActual.objects.get_or_create(
            usuario=mov.usuario,  # usuario dueño original del producto
            producto=mov.producto,
            defaults={"cantidad": 0},
        )

        # Sumar la cantidad eliminada
        stock_actual.cantidad += mov.cantidad_movida
        stock_actual.save()

        # Eliminar el movimiento
        mov.delete()

        return redirect("historial_restas_admin")


# ===============================================================================================
# ELIMINAR UNA RESTA - TECNICO
# ===============================================================================================


class EliminarRestaTecnicoView(LoginRequiredMixin, View):
    def get(self, request, pk):

        # Buscar movimiento SALIDA
        mov = Stock.objects.filter(id=pk, movimiento="SALIDA").first()
        if not mov:
            return redirect("historial_restas_admin")

        # Obtener o crear StockActual del usuario que tenía el producto
        stock_actual, creado = StockActual.objects.get_or_create(
            usuario=mov.usuario,  # usuario dueño original del producto
            producto=mov.producto,
            defaults={"cantidad": 0},
        )

        # Sumar la cantidad eliminada
        stock_actual.cantidad += mov.cantidad_movida
        stock_actual.save()

        # Eliminar el movimiento
        mov.delete()

        return redirect("mi_stock")


# ===============================================================================================
# STOCK POR USUARIO (ADMIN) - VER EL STOCK DE TODOS
# ===============================================================================================


class StockUsuariosAdministradorView(
    LoginRequiredMixin, StaffSuperuserRequiredMixin, FilterView
):
    model = StockActual
    template_name = "inventario/admin/stock_usuarios.html"
    context_object_name = "stocks"
    filterset_class = StockActualAdminFilter
    paginate_by = 20

    def get_queryset(self):
        # queryset base optimizado
        queryset = StockActual.objects.select_related("usuario", "producto").order_by(
            "usuario__email", "producto__nombre"
        )
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # mantener los parámetros en la URL durante la paginación
        query_params = self.request.GET.copy()
        if "page" in query_params:
            query_params.pop("page")

        context["querystring"] = query_params.urlencode()

        return context


@method_decorator(login_required, name="dispatch")
class ExportarStockExcelView(LoginRequiredMixin, View):
    def get(self, request, *args, **kwargs):

        # Aplicar filtros exactamente igual que en la vista original
        stock_filtrado = StockFilter(
            request.GET,
            queryset=Stock.objects.select_related(
                "usuario",
                "usuario_origen",
                "usuario_destino",
                "producto",
            ).all(),
        ).qs

        # Crear archivo Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Movimientos"

        # Encabezados
        columnas = [
            "Usuario",
            "Fecha",
            "Hora",
            "Origen",
            "Destino",
            "Cantidad",
            "Producto",
            "Movimiento",
            "Motivo",
            "Referencia",
            "Ticket",
        ]

        ws.append(columnas)

        # Agregar filas
        for item in stock_filtrado:
            ws.append(
                [
                    item.usuario.email,
                    item.created.strftime("%d/%m/%Y"),
                    item.created.strftime("%H:%M"),
                    item.usuario_origen.email if item.usuario_origen else "",
                    item.usuario_destino.email if item.usuario_destino else "",
                    item.cantidad_movida,
                    item.producto.nombre,
                    item.get_movimiento_display(),
                    item.get_motivo_display() or "",
                    item.motivo_personalizado or "",
                    item.codigo_ticket or "",
                ]
            )

        # Ajustar el ancho de cada columna
        for col_num, columna in enumerate(columnas, 1):
            col_letter = get_column_letter(col_num)
            ws.column_dimensions[col_letter].width = 20

        # Respuesta como archivo descargable
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = (
            'attachment; filename="movimientos_stock.xlsx"'
        )

        wb.save(response)
        return response
