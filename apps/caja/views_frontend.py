from django.shortcuts import redirect
from django.contrib import messages
from django.urls import reverse_lazy
from django.views.generic import (
    ListView,
    CreateView,
    DetailView,
    UpdateView,
    DeleteView,
)
from django.contrib.auth.mixins import LoginRequiredMixin
from decimal import Decimal

from .models import Caja, SaldoCaja
from .forms import CajaForm
from django.views import View
from django.shortcuts import redirect, get_object_or_404
from django.contrib import messages
from .models import SaldoCaja
from .filters import CajaFilter
from django_filters.views import FilterView

# views.py
import openpyxl
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.views import View
from django.contrib.auth.mixins import LoginRequiredMixin

from .filters import CajaFilter
from .models import Caja, SaldoCaja


class ExportarCajaExcelView(LoginRequiredMixin, View):
    def get(self, request, *args, **kwargs):

        # Aplicar filtros igual que en la vista principal
        movimientos = CajaFilter(
            request.GET,
            queryset=Caja.objects.filter(usuario=request.user).order_by("-created"),
        ).qs

        # Crear archivo Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Movimientos Caja"

        columnas = [
            "Fecha",
            "Hora",
            "Tipo Movimiento",
            "Motivo",
            "Cantidad Movida",
            "Saldo Actual",
            "Info Adicional",
        ]

        ws.append(columnas)

        # Obtener saldo actual del usuario
        saldo_caja = SaldoCaja.objects.filter(usuario=request.user).first()
        saldo_actual = saldo_caja.saldo if saldo_caja else "0.00"

        for mov in movimientos:
            ws.append(
                [
                    mov.created.strftime("%d/%m/%Y"),
                    mov.created.strftime("%H:%M"),
                    mov.get_tipo_movimiento_display(),
                    mov.get_motivo_movimiento_display(),
                    str(mov.cantidad_movida),
                    str(saldo_actual),
                    mov.info_adicional or "",
                ]
            )

        # Ajustar ancho de las columnas
        for col_num, columna in enumerate(columnas, 1):
            col_letter = get_column_letter(col_num)
            ws.column_dimensions[col_letter].width = 20

        # Respuesta HTTP con descarga
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="movimientos_caja.xlsx"'

        wb.save(response)
        return response


# =============================================================================
# FUNCION PARA OBTENER O CREAR SALDO
# =============================================================================
def obtener_saldo(usuario):
    saldo, creado = SaldoCaja.objects.get_or_create(
        usuario=usuario, defaults={"saldo": Decimal("0.00")}
    )
    return saldo


# =============================================================================
# LISTAR MOVIMIENTOS
# =============================================================================
class CajaListView(LoginRequiredMixin, FilterView):
    model = Caja
    template_name = "caja/movimientos_list.html"
    context_object_name = "movimientos"
    paginate_by = 20
    filterset_class = CajaFilter

    def get_queryset(self):
        # Filtra los movimientos del usuario logueado
        return Caja.objects.filter(usuario=self.request.user).order_by("created")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        movimientos = context.get("movimientos")
        saldo = Decimal("0.00")

        # Calcula el saldo resultante para cada movimiento filtrado
        for mov in movimientos:
            if mov.tipo_movimiento == "INGRESO":
                saldo += mov.cantidad_movida
            elif mov.tipo_movimiento == "SALIDA":
                saldo -= mov.cantidad_movida

            mov.saldo_resultante = saldo

        context["saldo_actual"] = obtener_saldo(self.request.user).saldo

        return context


# =============================================================================
# DETALLE
# =============================================================================
class CajaDetailView(LoginRequiredMixin, DetailView):
    model = Caja
    template_name = "caja/movimiento_detail.html"
    context_object_name = "movimiento"

    def get_queryset(self):
        return Caja.objects.filter(usuario=self.request.user)


# =============================================================================
# CREAR
# =============================================================================
class CajaCreateView(LoginRequiredMixin, CreateView):
    model = Caja
    form_class = CajaForm
    template_name = "caja/movimiento_create.html"
    success_url = reverse_lazy("caja_list")

    def form_valid(self, form):
        usuario = self.request.user
        saldo = obtener_saldo(usuario)

        form.instance.usuario = usuario
        movimiento = form.save(commit=False)
        cantidad = movimiento.cantidad_movida

        if movimiento.tipo_movimiento == "INGRESO":
            saldo.saldo += cantidad
        elif movimiento.tipo_movimiento == "SALIDA":
            if saldo.saldo < cantidad:
                messages.error(
                    self.request, "No hay saldo suficiente para esta salida."
                )
                return redirect("caja_create")
            saldo.saldo -= cantidad

        saldo.save()
        movimiento.save()

        messages.success(self.request, "Movimiento registrado correctamente.")
        return super().form_valid(form)


# =============================================================================
# EDITAR
# =============================================================================
class CajaUpdateView(LoginRequiredMixin, UpdateView):
    model = Caja
    form_class = CajaForm
    template_name = "caja/movimiento_update.html"
    success_url = reverse_lazy("caja_list")

    def get_queryset(self):
        return Caja.objects.filter(usuario=self.request.user)

    def form_valid(self, form):
        usuario = self.request.user
        saldo = obtener_saldo(usuario)

        movimiento_anterior = self.get_object()

        old_cantidad = movimiento_anterior.cantidad_movida
        old_tipo = movimiento_anterior.tipo_movimiento

        new_movimiento = form.save(commit=False)
        new_cantidad = new_movimiento.cantidad_movida
        new_tipo = new_movimiento.tipo_movimiento

        # Revertir el anterior
        if old_tipo == "INGRESO":
            saldo.saldo -= old_cantidad
        elif old_tipo == "SALIDA":
            saldo.saldo += old_cantidad

        # Aplicar el nuevo
        if new_tipo == "INGRESO":
            saldo.saldo += new_cantidad
        elif new_tipo == "SALIDA":
            if saldo.saldo < new_cantidad:
                messages.error(
                    self.request, "No hay saldo suficiente para este cambio."
                )
                return redirect("caja_update", pk=self.object.pk)
            saldo.saldo -= new_cantidad

        saldo.save()
        new_movimiento.save()

        messages.success(self.request, "Movimiento actualizado correctamente.")
        return super().form_valid(form)


# =============================================================================
# ELIMINAR
# =============================================================================
class CajaDeleteView(LoginRequiredMixin, DeleteView):
    model = Caja
    template_name = "caja/movimiento_confirm_delete.html"
    context_object_name = "movimiento"
    success_url = reverse_lazy("caja_list")

    def get_queryset(self):
        return Caja.objects.filter(usuario=self.request.user)

    def delete(self, request, *args, **kwargs):
        movimiento = self.get_object()
        saldo = obtener_saldo(request.user)

        # REVERSIÓN CORRECTA
        if movimiento.tipo_movimiento == "INGRESO":
            saldo.saldo -= movimiento.cantidad_movida
        elif movimiento.tipo_movimiento == "SALIDA":
            saldo.saldo += movimiento.cantidad_movida

        saldo.save()

        messages.success(
            request, "Movimiento eliminado y saldo ajustado correctamente."
        )
        return super().delete(request, *args, **kwargs)


class SaldoResetView(LoginRequiredMixin, View):
    def post(self, request, user_id):
        # Solo permitir que el usuario resetee SU propio saldo
        if str(request.user.id) != str(user_id):
            messages.error(request, "No tienes permiso para realizar esta acción.")
            return redirect("caja_list")

        try:
            saldo = SaldoCaja.objects.get(usuario=request.user)
        except SaldoCaja.DoesNotExist:
            messages.error(request, "No se encontró saldo asociado al usuario.")
            return redirect("caja_list")

        # Resetear el saldo
        saldo.saldo = 0
        saldo.save()

        messages.success(request, "Saldo reseteado correctamente.")
        return redirect("caja_list")
